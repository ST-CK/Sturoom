from fastapi import APIRouter, Request
import requests, os
from openai import OpenAI
from dotenv import load_dotenv
from PyPDF2 import PdfReader
from pptx import Presentation
from io import BytesIO

load_dotenv()
client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
router = APIRouter()

@router.post("/from-url")
async def generate_quiz_from_url(req: Request):
    data = await req.json()
    file_url = data.get("file_url")
    mode = data.get("mode", "multiple")

    # 1️⃣ 파일 다운로드
    res = requests.get(file_url)
    if res.status_code != 200:
        return {"error": "파일 다운로드 실패"}

    filename = file_url.split("/")[-1].lower()

    # 2️⃣ 파일 내용 읽기
    text = ""
    if filename.endswith(".pdf"):
        reader = PdfReader(BytesIO(res.content))
        for page in reader.pages:
            text += page.extract_text() + "\n"
    elif filename.endswith(".pptx"):
        prs = Presentation(BytesIO(res.content))
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
    else:
        return {"error": "지원하지 않는 파일 형식"}

    # 3️⃣ GPT로 퀴즈 생성
    prompt = f"""
    다음은 강의 자료입니다:
    ---
    {text[:4000]}  # 너무 긴 경우 앞부분만 사용
    ---
    위 내용을 바탕으로 {mode} 문제 1개를 JSON으로 만들어줘.
    """

    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": prompt}],
    )

    quiz_text = response.choices[0].message.content
    return {"quiz": quiz_text}
